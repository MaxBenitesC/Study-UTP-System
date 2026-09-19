#!/usr/bin/env python3
"""Genera los cuatro entregables de la semana 5 desde el proyecto de semana 4."""

from copy import copy
from datetime import date, timedelta
from pathlib import Path
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE = Path(__file__).resolve().parent
OUT = BASE / "ENTREGABLES"
OUT.mkdir(exist_ok=True)
PROJECT = "Modernización del Sistema de Punto de Venta (POS) e Implementación de Cajas de Autoservicio (Self-Checkout) en Tiendas Retail"
TEAM = "Equipo N.° 10 — Benites Corazón, Max Anderson (U24217839); Oliva Seminario, Anderson Harold (U21222001)"
START = date(2026, 9, 14)
NAVY = "17324D"
TEAL = "087E8B"
PALE = "EAF2F6"


# Categoría, ID, requisito, responsable, relación, comprobación, paquete EDT.
REQUIREMENTS = [
    ("Negocio", "RN01", "Reducir el tiempo de espera en caja frente a la situación inicial descrita en el acta (>12 minutos en horas punta), midiendo el antes y después en la tienda piloto.", "Patrocinador / Operaciones", "Acta de constitución", "Medición comparativa de tiempos de espera en horas punta.", "4.3 / 6.1"),
    ("Negocio", "RN02", "Poner en operación las estaciones de autoservicio y actualizar las cajas POS previstas en el proyecto dentro de 12 semanas.", "Patrocinador", "Acta de constitución", "Acta de entrega y cronograma final de 12 semanas.", "5.3 / 6.2"),
    ("Negocio", "RN03", "Mantener el costo del proyecto dentro del presupuesto de S/ 200 000 indicado en el acta, con desviación inferior al 5 % respecto del presupuesto aprobado.", "Patrocinador / Finanzas", "Acta de constitución", "Reporte final de costos y comparación con línea base.", "1.4 / 6.3"),
    ("Negocio", "RN04", "Realizar las instalaciones sin suspender la atención normal de la tienda durante el horario comercial.", "Gerencia de tienda", "Acta de constitución", "Bitácora de ventanas de trabajo y registro de continuidad operativa.", "2.3 / 5.3"),
    ("Interesados", "RI01", "El gerente de tienda debe aprobar las ventanas nocturnas de instalación y recibir aviso previo de cambios que afecten cajas operativas.", "Gerencia de tienda", "RN04", "Cronograma de intervención validado por la tienda.", "1.4 / 2.3"),
    ("Interesados", "RI02", "El personal de caja y los supervisores deben recibir capacitación práctica para operar y asistir el autoservicio.", "Supervisores de caja", "RN02", "Lista de asistencia y demostración práctica registrada.", "5.2"),
    ("Interesados", "RI03", "El área de TI debe validar conectividad, configuración de POS y respaldo antes de la apertura al público.", "Área de TI", "RN02", "Lista de comprobación técnica firmada.", "2.3 / 3.1 / 4.2"),
    ("Interesados", "RI04", "El patrocinador y la gerencia de tienda deben revisar los resultados de las pruebas y aprobar la aceptación final.", "Patrocinador / Gerencia de tienda", "RN01 / RN02", "Acta de aceptación con evidencias de prueba.", "6.2"),
    ("Interesados", "RI05", "El proveedor de equipos y pagos debe entregar componentes, soporte y documentación de garantía según las especificaciones aprobadas.", "Proveedor", "RN02", "Guías de recepción y garantías documentadas.", "2.2 / 5.1"),
    ("Solución", "RS01", "Cada estación de autoservicio debe permitir lectura de productos, visualización de la compra y pago con tarjeta.", "TI / Operaciones", "RI03 / RI04", "Prueba funcional completa por cada estación instalada.", "3.2 / 4.2"),
    ("Solución", "RS02", "Las cajas POS contempladas deben operar con el software actualizado y conservar el flujo de cobro definido en el proyecto.", "TI / Cajas", "RI03", "Pruebas funcionales y conformidad del supervisor.", "3.1 / 4.2"),
    ("Solución", "RS03", "La solución debe conectarse a la pasarela de pagos con tarjeta y registrar transacciones aprobadas y rechazadas correctamente.", "TI / Proveedor de pagos", "RI03", "Casos de prueba de integración con transacciones controladas.", "3.3 / 4.2"),
    ("Solución", "RS04", "Cada punto de cobro debe disponer de red y energía estabilizada verificadas antes de instalar equipos.", "TI / Infraestructura", "RI03", "Acta de pruebas de red y energía por punto.", "2.3"),
    ("Solución", "RS05", "El proceso de autoservicio debe lograr un tiempo promedio de atención inferior a dos minutos por cliente durante la prueba piloto.", "Operaciones", "RN01", "Muestra cronometrada del piloto y promedio calculado.", "4.3"),
    ("Solución", "RS06", "Los terminales deben emitir el comprobante de venta conforme al flujo de facturación electrónica previsto por el proyecto.", "TI / Finanzas", "RI03", "Prueba controlada de emisión y validación del comprobante.", "3.3 / 4.2"),
    ("Transición", "RT01", "Antes de la apertura se debe ejecutar un piloto con compras controladas y registrar incidencias y correcciones.", "Operaciones / TI", "RS01 / RS03", "Informe de piloto con resultados e incidencias resueltas.", "4.3"),
    ("Transición", "RT02", "La puesta en operación debe incluir un procedimiento de reversión y soporte inicial ante fallas de cobro o conectividad.", "TI / Gerencia de tienda", "RN04 / RS03", "Plan de reversión aprobado y registro de soporte inicial.", "5.3 / 6.1"),
    ("Transición", "RT03", "Los supervisores deben recibir manuales de uso y atención de incidencias antes de la apertura.", "Supervisores", "RI02", "Entrega documentada de manuales y capacitación.", "5.1 / 5.2"),
    ("Proyecto", "RP01", "El plan debe contener requisitos, enunciado del alcance, EDT y cronograma consistentes con el acta de constitución.", "Gerente de proyecto", "Acta de constitución", "Revisión cruzada y aprobación de documentos de planificación.", "1.2 / 1.3 / 1.4"),
    ("Proyecto", "RP02", "Los hitos del proyecto deben respetar la secuencia del acta: plan y orden de compra en semana 3, equipos e infraestructura en semana 6, piloto en semana 9, apertura en semana 11 y cierre en semana 12.", "Gerente de proyecto", "Acta de constitución", "Cronograma con hitos y fechas correspondientes.", "1.4 / 6.3"),
    ("Proyecto", "RP03", "Los cambios de alcance, costo o plazo deben registrarse y someterse a aprobación del patrocinador antes de incorporarlos.", "Gerente de proyecto / Patrocinador", "RN02 / RN03", "Registro de cambios y decisiones de aprobación.", "1.4 / 6.3"),
    ("Calidad", "RC01", "Todas las cajas instaladas deben completar las pruebas de cobro previstas sin fallas antes de la aceptación.", "TI / Operaciones", "RS01 / RS02 / RS03", "Matriz de pruebas aprobada y acta de resultados.", "4.2 / 6.2"),
    ("Calidad", "RC02", "La aceptación final requiere funcionamiento correcto de las cajas instaladas durante una semana continua.", "Patrocinador / Gerencia de tienda", "RI04", "Bitácora de operación de una semana y acta de aceptación.", "6.1 / 6.2"),
    ("Calidad", "RC03", "La entrega final debe incluir personal capacitado, manuales y documentos de garantía.", "Gerencia de tienda", "RI02 / RI05", "Acta de entrega de manuales, garantías y capacitación.", "5.1 / 5.2 / 6.2"),
]


# ID, paquete orientado a entregable, semana inicial/final, responsable, actividades, precedencias.
PACKAGES = [
    ("1.1", "Acta y registro de interesados", 1, 1, "Gerente de proyecto", "Revisar el acta y validar interesados", "Formalizar aprobación del acta y registro", ""),
    ("1.2", "Matriz de requisitos", 1, 2, "Gerente de proyecto / usuarios", "Recoger necesidades de operación, TI y cajas", "Clasificar y validar requisitos por categoría", "1.1"),
    ("1.3", "Enunciado del alcance y EDT", 2, 2, "Gerente de proyecto", "Definir entregables, aceptación y exclusiones", "Descomponer todos los entregables en la EDT", "1.2"),
    ("1.4", "Cronograma y plan aprobados", 3, 3, "Gerente de proyecto / patrocinador", "Estimar actividades, secuencia y costos", "Validar plan, línea base y orden de compra", "1.3"),
    ("2.1", "Orden de compra de equipos", 3, 3, "Compras / patrocinador", "Confirmar especificaciones y cotizaciones", "Emitir y aprobar orden de compra", "1.4"),
    ("2.2", "Equipos recibidos y verificados", 6, 6, "Compras / TI", "Recepcionar terminales POS y autoservicio", "Inspeccionar equipos y registrar garantías", "2.1"),
    ("2.3", "Puntos de red y energía habilitados", 4, 6, "TI / infraestructura", "Adecuar red y puntos eléctricos en ventanas autorizadas", "Probar conectividad y energía estabilizada", "1.4"),
    ("3.1", "Software POS actualizado", 7, 8, "TI / proveedor POS", "Preparar versión y respaldo del software POS", "Actualizar cajas previstas y verificar arranque", "2.2, 2.3"),
    ("3.2", "Estaciones de autoservicio instaladas", 7, 9, "TI / proveedor SCO", "Instalar muebles, terminales y lectores", "Configurar autoservicio y validar dispositivos", "2.2, 2.3"),
    ("3.3", "Pagos y comprobantes integrados", 8, 9, "TI / proveedor de pagos", "Configurar pasarela y emisión de comprobantes", "Validar transacciones controladas de integración", "3.1, 3.2"),
    ("4.1", "Plan y casos de prueba", 8, 8, "TI / operaciones", "Redactar casos de cobro, red y contingencia", "Aprobar plan de pruebas y criterios de aceptación", "3.1"),
    ("4.2", "Pruebas funcionales y de integración", 9, 9, "TI / supervisores", "Ejecutar pruebas por caja y registrar resultados", "Corregir incidencias y repetir pruebas fallidas", "3.3, 4.1"),
    ("4.3", "Piloto y acta de resultados", 9, 10, "Operaciones / TI", "Ejecutar piloto con compras controladas", "Medir tiempos y cerrar incidencias del piloto", "4.2"),
    ("5.1", "Manuales y garantías entregados", 10, 10, "TI / proveedor", "Preparar manuales de uso y soporte", "Reunir garantías y entregar documentación", "4.3"),
    ("5.2", "Personal capacitado", 11, 11, "Operaciones / supervisores", "Capacitar a supervisores y cajeros", "Evaluar operación práctica y registrar asistencia", "5.1"),
    ("5.3", "Apertura de cajas al público", 11, 11, "Gerencia de tienda / TI", "Autorizar puesta en marcha y verificar reversión", "Abrir servicio con soporte en tienda", "4.3, 5.2"),
    ("6.1", "Bitácora de operación inicial", 11, 12, "Gerencia de tienda / TI", "Monitorear operación continua durante una semana", "Consolidar incidencias y métricas de servicio", "5.3"),
    ("6.2", "Acta de aceptación", 12, 12, "Patrocinador / gerencia de tienda", "Revisar pruebas, métricas y entregables", "Firmar aceptación formal del proyecto", "6.1"),
    ("6.3", "Informe de cierre", 12, 12, "Gerente de proyecto", "Conciliar costos, cronograma y lecciones aprendidas", "Entregar informe final y cerrar el proyecto", "6.2"),
]
BRANCHES = [
    ("1", "Gestión y planificación"),
    ("2", "Suministro e infraestructura"),
    ("3", "Configuración e integración"),
    ("4", "Pruebas y piloto"),
    ("5", "Capacitación y apertura"),
    ("6", "Seguimiento y cierre"),
]


def setup_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor(23, 50, 77)
    footer = sec.footer.paragraphs[0]
    footer.text = "Gestión de Proyectos · Semana 05 · Equipo N.° 10"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.author = "Equipo N.° 10"
    doc.core_properties.last_modified_by = "Equipo N.° 10"


def add_cover(doc):
    """Replica la disposición de la carátula de TA_Gestion_de_Proyectos.docx."""
    logo = BASE.parent / "SEMANA04" / "logo_utp.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    p.add_run().add_picture(str(logo), width=Inches(2.8))

    def line(value, size=11, bold=False, italic=False, after=3):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(value)
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        return p

    line("FACULTAD DE INGENIERÍA", 14, True, after=28)
    line("TAREA ACADÉMICA", 16, True, after=28)
    line("Semana 05: Elaboración del EDT y enunciado del alcance del proyecto", 12, italic=True, after=22)
    line("CURSO:", 11, True, after=2)
    line("GESTIÓN DE PROYECTOS (100000I33N) - Sección 13451", 12, True, after=18)
    line("DOCENTE:", 11, True, after=2)
    line("Nora Noelia Sernaque Zapata", 11, after=18)
    line("INTEGRANTES:", 11, True, after=6)
    line("Benites Corazón, Max Anderson – Código: U24217839", 11, after=3)
    line("Oliva Seminario, Anderson Harold – Código: U21222001", 11, after=3)
    p = line("Lima - Perú\n2026", 11, after=0)
    p.paragraph_format.space_before = Pt(20)
    doc.add_page_break()


def add_header(doc, subtitle):
    add_cover(doc)
    doc.add_heading("Gestión de Proyectos — Semana 05", 0)
    doc.add_paragraph(subtitle, style="Subtitle")
    doc.add_paragraph(PROJECT)
    doc.add_paragraph(TEAM)
    doc.add_paragraph("Base: acta de constitución y registro de interesados de la semana 04. Versión académica propuesta para revisión del equipo.")


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = 1
            if i == 0:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                cell._tc.get_or_add_tcPr().append(_shade(NAVY))


def _shade(color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    sh = OxmlElement("w:shd")
    sh.set(qn("w:fill"), color)
    return sh


def make_requirements():
    src = BASE / "MATERIALES" / "Formato_Requerimientos_por_Categoria_Tarea.xlsx"
    wb = load_workbook(src)
    ws = wb.active
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    for row in ws:
        for cell in row:
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
    ws.title = "Requerimientos"
    ws.merge_cells("A1:G1")
    ws["A1"] = "MATRIZ DE REQUERIMIENTOS — PROYECTO POS Y AUTOSERVICIO"
    ws["A2"] = "Proyecto"
    ws.merge_cells("B2:G2")
    ws["B2"] = PROJECT
    ws["A3"] = "Equipo"
    ws.merge_cells("B3:G3")
    ws["B3"] = TEAM
    ws["A4"] = "Base"
    ws.merge_cells("B4:G4")
    ws["B4"] = "Acta de constitución de semana 04; seis categorías de requisitos indicadas por UTP."
    headers = ["Categoría", "ID", "Requisito verificable", "Interesado / responsable", "Relación", "Verificación / aceptación", "Paquete EDT"]
    for col, val in enumerate(headers, 1):
        c = ws.cell(6, col, val)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rownum, values in enumerate(REQUIREMENTS, 7):
        for col, val in enumerate(values, 1):
            c = ws.cell(rownum, col, val)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=Side(style="hair", color="CBD5DF"))
            if rownum % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F6F9")
        ws.row_dimensions[rownum].height = 52
    widths = [17, 10, 78, 29, 22, 48, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 48
    ws.row_dimensions[6].height = 34
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", bold=True, color="FFFFFF", size=15)
    ws["A1"].alignment = Alignment(vertical="center")
    for r in (2, 3, 4):
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "C7"
    ws.auto_filter.ref = f"A6:G{6 + len(REQUIREMENTS)}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:G{6 + len(REQUIREMENTS)}"
    wb.save(OUT / "01_Requerimientos_por_Categoria.xlsx")


def make_scope():
    doc = Document()
    setup_doc(doc)
    add_header(doc, "Entregable 2 de 4 — Enunciado del alcance del proyecto")
    doc.add_heading("1. Descripción del alcance del producto", 1)
    doc.add_paragraph("El proyecto comprende la modernización de las cajas POS previstas en el acta y la incorporación de estaciones de autoservicio en las tiendas consideradas por el patrocinador. El resultado será un conjunto de puntos de cobro instalados, conectados y probados, con lectura de productos, cobro con tarjeta y emisión de comprobantes, acompañado de personal capacitado y soporte de arranque. La cantidad exacta de cajas y tiendas se fijará en el inventario aprobado antes de la compra; el acta de semana 04 no define cantidades.")
    doc.add_paragraph("La finalidad es reducir las esperas frente a la situación inicial descrita en el acta, sin interrumpir la atención comercial. La meta de calidad del acta es un tiempo promedio de atención en autoservicio inferior a dos minutos por cliente durante el piloto. El proyecto se planifica para 12 semanas y un presupuesto de S/ 200 000, sujeto a la aprobación del patrocinador.")
    doc.add_heading("2. Entregables y criterios de aceptación", 1)
    rows = [
        ("Plan de proyecto", "Acta, interesados, requisitos, alcance, EDT, cronograma y control de cambios revisados por el patrocinador."),
        ("Equipos y puntos preparados", "Equipos recepcionados; garantías documentadas; red y energía estabilizada verificadas por punto."),
        ("POS y autoservicio configurados", "Cajas POS actualizadas y estaciones SCO instaladas, con prueba de lectura, cobro y emisión de comprobantes."),
        ("Integración de pagos", "Casos de pago aprobado y rechazado ejecutados y documentados sin fallas de cobro."),
        ("Piloto operativo", "Compras controladas, incidencias resueltas y promedio de atención en SCO inferior a dos minutos por cliente."),
        ("Capacitación y documentación", "Supervisores y cajeros capacitados; manuales y garantías entregados."),
        ("Aceptación y cierre", "Una semana continua de funcionamiento correcto de las cajas instaladas, acta firmada por patrocinador y gerencia de tienda, informe final de costo y plazo."),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Entregable", "Criterio de aceptación"
    for a, b in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = a, b
    style_table(t)
    doc.add_heading("3. Exclusiones", 1)
    for item in [
        "No se incluye renovación de todas las cajas de la cadena; solo las aprobadas en el inventario del proyecto.",
        "No se incluye desarrollo de una nueva pasarela de pagos ni modificación de sistemas bancarios externos.",
        "No se incluye atención permanente posterior al período de soporte inicial y aceptación, salvo contrato separado.",
        "No se incluyen obras civiles mayores ni remodelación integral de la tienda fuera de los puntos necesarios para POS y autoservicio.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("4. Supuestos y restricciones", 1)
    doc.add_paragraph("Supuestos: espacio disponible en tienda, entrega oportuna de equipos, capacidad de red e internet suficiente y participación del personal en la capacitación, tal como recoge el acta de constitución.")
    doc.add_paragraph("Restricciones: 12 semanas de duración; presupuesto base de S/ 200 000; instalaciones sin suspender la atención en horario comercial; cumplimiento de las condiciones aplicables a pagos con tarjeta y comprobantes electrónicos contempladas en el acta.")
    doc.add_heading("5. Trazabilidad y control", 1)
    doc.add_paragraph("La matriz de requisitos identifica cada necesidad y su paquete EDT. La EDT reúne el 100 % de los entregables aquí definidos y el cronograma descompone todos sus paquetes en actividades. Cualquier cambio de alcance, costo o plazo deberá registrarse y someterse a aprobación del patrocinador antes de incorporarlo.")
    doc.add_heading("6. Aprobación propuesta", 1)
    doc.add_paragraph("Patrocinador: Dirección de Operaciones y Finanzas. Conformidad operativa: gerencia de tienda. Este documento es una propuesta académica y requiere validación de cantidades, fechas y línea base antes de su ejecución real.")
    doc.save(OUT / "02_Enunciado_del_Alcance.docx")


def make_edt():
    dot = [
        "digraph EDT {",
        'graph [rankdir=TB, nodesep=0.45, ranksep=0.45, bgcolor="white"];',
        'node [shape=box, style="rounded,filled", color="#087E8B", fillcolor="#EAF2F6", fontname="Arial", fontsize=10];',
        'edge [color="#7B91A5", arrowsize=0.5];',
        'root [label="Proyecto POS y autoservicio", fillcolor="#17324D", fontcolor="white", fontsize=14];',
    ]
    for code, title in BRANCHES:
        dot.append(f'b{code} [label="{code}. {title}"]; root -> b{code};')
    dot.append("}")
    subprocess.run(["dot", "-Tpng", "-Gdpi=160", "-o", str(OUT / "EDT_Resumen.png")], input="\n".join(dot), text=True, check=True)
    doc = Document()
    setup_doc(doc)
    add_header(doc, "Entregable 3 de 4 — Estructura de Desglose del Trabajo (EDT)")
    doc.add_paragraph("EDT orientada a entregables. El nivel 1 representa el proyecto; el nivel 2, sus seis conjuntos de entregables; y el nivel 3, los paquetes de trabajo completos. Se incluyen tanto los entregables del producto como los de gestión, seguimiento y cierre.")
    doc.add_picture(str(OUT / "EDT_Resumen.png"), width=Cm(16))
    doc.add_heading("Desglose completo de paquetes de trabajo", 1)
    for code, title in BRANCHES:
        doc.add_heading(f"{code}. {title}", 2)
        t = doc.add_table(rows=1, cols=3)
        t.rows[0].cells[0].text = "Código"
        t.rows[0].cells[1].text = "Paquete de trabajo / entregable"
        t.rows[0].cells[2].text = "Semana prevista"
        for p in PACKAGES:
            if p[0].split(".")[0] != code:
                continue
            cells = t.add_row().cells
            cells[0].text = p[0]
            cells[1].text = p[1]
            cells[2].text = str(p[2]) if p[2] == p[3] else f"{p[2]}–{p[3]}"
        style_table(t)
    doc.add_paragraph("Regla del 100 %: los paquetes anteriores cubren planificación, suministro, infraestructura, configuración, pruebas, capacitación, apertura, seguimiento y cierre establecidos en el enunciado del alcance. Las actividades y fechas de cada paquete aparecen en el cronograma adjunto.")
    doc.save(OUT / "03_EDT_Proyecto_POS_SCO.docx")


def week_date(week, day):
    return START + timedelta(days=(week - 1) * 7 + day)


def make_schedule():
    src = BASE / "MATERIALES" / "Formato_Cronograma_Tarea.xlsx"
    wb = load_workbook(src)
    ws = wb.active
    for row in ws:
        for cell in row:
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
    ws.title = "Cronograma EDT"
    ws.merge_cells("A1:J1")
    ws["A1"] = "CRONOGRAMA DEL PROYECTO POS Y AUTOSERVICIO — 12 SEMANAS"
    ws.merge_cells("A2:J2")
    ws["A2"] = PROJECT
    ws.merge_cells("A3:J3")
    ws["A3"] = "Línea base académica propuesta: 14/09/2026–06/12/2026. Fechas referenciales sujetas a aprobación del patrocinador."
    headers = ["Código", "Paquete EDT", "Actividad", "Responsable", "Inicio", "Fin", "Días calendario", "Predecesora(s)", "Semana(s)", "Hito / evidencia"]
    for col, val in enumerate(headers, 1):
        c = ws.cell(5, col, val)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row = 6
    package_end = {p[0]: week_date(p[3], 4) for p in PACKAGES}
    for code, title, w1, w2, owner, action1, action2, depends in PACKAGES:
        first_start = week_date(w1, 0)
        first_end = week_date(w1, 2)
        second_start = week_date(w1, 3)
        second_end = week_date(w2, 4)
        if code == "5.2":
            first_start, first_end = week_date(11, 0), week_date(11, 1)
            second_start, second_end = week_date(11, 2), week_date(11, 3)
        elif code == "5.3":
            first_start = first_end = week_date(11, 3)
            second_start = second_end = week_date(11, 4)
        elif code == "6.1":
            first_start, first_end = week_date(11, 4), week_date(12, 3)
            second_start = second_end = week_date(12, 3)
        elif code == "6.2":
            first_start = first_end = week_date(12, 3)
            second_start = second_end = week_date(12, 4)
        elif code == "6.3":
            first_start, first_end = week_date(12, 1), week_date(12, 3)
            second_start = second_end = week_date(12, 4)
        valid_deps = [x.strip() for x in depends.split(",") if x.strip() and package_end[x.strip()] <= first_start]
        prev = ", ".join(x + ".2" for x in valid_deps)
        for actcode, action, start, end, pred in [
            (f"{code}.1", action1, first_start, first_end, prev),
            (f"{code}.2", action2, second_start, second_end, f"{code}.1"),
        ]:
            vals = [actcode, f"{code} {title}", action, owner, start, end, None, pred, str(w1) if w1 == w2 else f"{w1}–{w2}", ""]
            for col, val in enumerate(vals, 1):
                c = ws.cell(row, col, val)
                c.font = Font(name="Arial", size=9)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = Border(bottom=Side(style="hair", color="CBD5DF"))
                if row % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F2F6F9")
            ws.cell(row, 7, f"=F{row}-E{row}+1")
            ws.cell(row, 5).number_format = "dd/mm/yyyy"
            ws.cell(row, 6).number_format = "dd/mm/yyyy"
            ws.row_dimensions[row].height = 38
            row += 1
    milestones = [
        (1, "Acta aprobada", "1.1.2"),
        (3, "Plan y orden de compra aprobados", "2.1.2"),
        (6, "Equipos y puntos habilitados", "2.2.2, 2.3.2"),
        (9, "Instalación y piloto iniciados", "3.2.2, 4.3.1"),
        (11, "Personal capacitado y apertura", "5.2.2, 5.3.2"),
        (12, "Aceptación y cierre", "6.2.2, 6.3.2"),
    ]
    ws.cell(row + 1, 1, "HITOS DEL ACTA DE CONSTITUCIÓN")
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=10)
    ws.cell(row + 1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(row + 1, 1).font = Font(name="Arial", bold=True, color="FFFFFF")
    for n, (week, title, links) in enumerate(milestones, row + 2):
        ws.cell(n, 1, f"H{week}")
        ws.cell(n, 2, title)
        ws.cell(n, 5, week_date(week, 4))
        ws.cell(n, 5).number_format = "dd/mm/yyyy"
        ws.cell(n, 8, links)
        ws.cell(n, 9, str(week))
        ws.row_dimensions[n].height = 25
    widths = [13, 37, 51, 28, 16, 16, 14, 21, 13, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", bold=True, color="FFFFFF", size=15)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 45
    ws.row_dimensions[3].height = 29
    ws.row_dimensions[5].height = 35
    for r in (2, 3):
        ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:J{row - 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:J{row + 1 + len(milestones)}"
    wb.save(OUT / "04_Cronograma_basado_en_EDT.xlsx")


if __name__ == "__main__":
    make_requirements()
    make_scope()
    make_edt()
    make_schedule()
    print("Generados los cuatro entregables editables en", OUT)
