import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
import zipfile

def populate_acta_constitucion():
    src_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/MATERIALES/Formato_Acta_de_Constitucion.xlsx'
    out_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/Acta_de_Constitucion.xlsx'
    
    wb = openpyxl.load_workbook(src_path)
    ws = wb.active

    font_label = Font(name='Arial', size=10, bold=True, color='000000')
    font_content = Font(name='Arial', size=10, bold=False, color='000000')
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    data = {
        6: (
            "Equipo N° 10 - Integrantes:\n"
            "• Benites Corazón, Max Anderson (Código: U24217839)\n"
            "• Oliva Seminario, Anderson Harold (Código: U21222001)"
        ),
        7: (
            "Modernización del Sistema de Punto de Venta (POS) e Implementación de Cajas de Autoservicio (Self-Checkout) en Tiendas Retail"
        ),
        8: (
            "Nombre: Max Anderson Benites Corazón\n"
            "Cargo: Gerente de Proyecto (Project Manager)\n"
            "Organización: Equipo de Proyectos de Implementación Tecnológica\n"
            "Nivel de Autoridad: Autoridad para coordinar las actividades del equipo, supervisar el cronograma y entregables, "
            "coordinar con proveedores y reportar el avance formal al patrocinador.\n"
            "A quién responde: Responde directamente al Patrocinador del Proyecto (Dirección de Operaciones)."
        ),
        9: (
            "Patrocinador: Dirección de Operaciones y Finanzas de la Empresa\n"
            "Nivel de Autoridad: Máxima autoridad del proyecto; asigna el presupuesto, aprueba formalmente el Acta de Constitución "
            "y autoriza el paso a las fases de ejecución y cierre."
        ),
        10: (
            "El proyecto consiste en la renovación tecnológica de las cajas de cobro tradicionales (POS) y la incorporación de módulos "
            "de autoservicio (Self-Checkout) en las tiendas de la cadena.\n"
            "• Producto final: Estaciones de cobro modernas y estaciones de autoservicio operativas con lectores ópticos y pago con tarjeta.\n"
            "• Necesidad que soluciona: Disminuir las largas colas en horarios de mayor afluencia y agilizar el proceso de compra para clientes con pocos productos.\n"
            "• Alineamiento estratégico: Apoya el objetivo del negocio de mejorar la experiencia del cliente y modernizar la infraestructura comercial."
        ),
        11: (
            "1. Reducción de tiempos de espera: En horas punta los clientes esperan más de 12 minutos en caja, lo cual causa malestar y pérdida de ventas.\n"
            "2. Eficiencia en tienda: Las cajas de autoservicio permiten atender hasta 3 veces más clientes por metro cuadrado de caja.\n"
            "3. Tendencia del mercado: Los clientes demandan cada vez más opciones de pago rápido sin contacto y autoservicio."
        ),
        12: (
            "• Objetivo de Alcance: Instalar y poner en funcionamiento las cajas de autoservicio y actualizar el software en las cajas tradicionales programadas.\n"
            "  - Criterio de éxito: Acta de entrega y pruebas de funcionamiento al 100% sin fallas de cobro.\n"
            "• Objetivo de Tiempo: Culminar la implementación en un plazo de 12 semanas.\n"
            "  - Criterio de éxito: Cumplir con la fecha límite sin retrasos en la atención diaria de la tienda.\n"
            "• Objetivo de Costo: No exceder el presupuesto estimado para equipos y servicios.\n"
            "  - Criterio de éxito: Desviación menor al 5% del presupuesto aprobado.\n"
            "• Objetivo de Calidad: Lograr que el tiempo promedio de atención por cliente en autoservicio sea menor a 2 minutos."
        ),
        13: (
            "1. Adquisición e instalación física de los muebles y terminales de autoservicio.\n"
            "2. Conexión de cableado de red y puntos eléctricos estabilizados para cada caja.\n"
            "3. Instalación y configuración del software de punto de venta y conexión con la pasarela de pagos con tarjeta.\n"
            "4. Pruebas de funcionamiento y pruebas piloto con compras reales controladas.\n"
            "5. Capacitación a los supervisores y personal de caja para asistir a los clientes en las nuevas estaciones."
        ),
        14: (
            "1. Retraso en la llegada o importación de los equipos de autoservicio (Mitigación: Realizar los pedidos con suficiente anticipación y seguimiento continuo al proveedor).\n"
            "2. Fallas en la red o conexión a internet de la tienda durante la instalación (Mitigación: Contar con enlace de contingencia y verificar el cableado previamente).\n"
            "3. Resistencia o confusión inicial de los clientes al usar el nuevo sistema (Mitigación: Asignar personal de apoyo y señalización clara en los primeros días).\n"
            "4. Dificultad para coordinar los trabajos sin interrumpir la atención de la tienda (Mitigación: Programar las instalaciones en horarios nocturnos o fuera de atención al público)."
        ),
        15: (
            "• Hito 1 (Semana 1): Aprobación del Acta de Constitución del Proyecto.\n"
            "• Hito 2 (Semana 3): Aprobación del plan de trabajo y orden de compra de equipos.\n"
            "• Hito 3 (Semana 6): Llegada de equipos y adecuación de puntos de red y energía.\n"
            "• Hito 4 (Semana 9): Instalación, configuración y pruebas piloto en tienda.\n"
            "• Hito 5 (Semana 11): Capacitación del personal de tienda y apertura al público.\n"
            "• Hito 6 (Semana 12): Evaluación de resultados, aceptación formal y cierre del proyecto."
        ),
        16: (
            "Presupuesto estimado de alto nivel (orden de magnitud):\n"
            "• Equipamiento de autoservicio y terminales POS: S/ 120,000\n"
            "• Materiales de red, cableado y adecuación eléctrica: S/ 25,000\n"
            "• Servicios de configuración, instalación y soporte técnico: S/ 35,000\n"
            "• Capacitación y material de difusión en tienda: S/ 10,000\n"
            "• Fondo de contingencia para imprevistos: S/ 10,000\n"
            "Total estimado: S/ 200,000"
        ),
        17: (
            "1. La tienda contará con el espacio físico disponible para instalar las nuevas estaciones.\n"
            "2. Los proveedores entregarán los equipos en las fechas acordadas en la orden de compra.\n"
            "3. La conexión de red e internet de la tienda soportará las transacciones de las nuevas cajas.\n"
            "4. El personal de tienda participará en las capacitaciones programadas."
        ),
        18: (
            "1. Los trabajos de instalación física no deben suspender la atención normal de los clientes en horario comercial.\n"
            "2. El proyecto debe ajustarse estrictamente al presupuesto aprobado por la dirección.\n"
            "3. El sistema de cobro debe cumplir con las normas de seguridad de tarjetas y comprobantes electrónicos de SUNAT."
        ),
        19: (
            "• Dirección de la Empresa y Patrocinador (Área que financia y aprueba).\n"
            "• Gerente de Tienda y Supervisores (Responsables de la operación diaria).\n"
            "• Personal de Cajas / Anfitriones (Usuarios operativos que guían a los clientes).\n"
            "• Clientes de la tienda (Usuarios finales que realizan el pago de sus compras).\n"
            "• Área de TI y Sistemas (Soporte de red, software y equipos).\n"
            "• Proveedor de equipos y pasarela de pago (Suministro técnico)."
        ),
        20: (
            "• Quién aprueba: El Patrocinador del Proyecto (Dirección de Operaciones) junto con el Gerente de Tienda.\n"
            "• Criterios de aprobación:\n"
            "  - Funcionamiento correcto de todas las cajas instaladas durante una semana continua.\n"
            "  - Personal capacitado y conforme con el manejo del sistema.\n"
            "  - Entrega de manuales y documentos de garantía de los equipos."
        )
    }

    for row_idx, content in data.items():
        cell = ws.cell(row=row_idx, column=2)
        cell.value = content
        cell.font = font_content
        cell.alignment = align_left
        ws.cell(row=row_idx, column=1).font = font_label
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        ws.cell(row=row_idx, column=1).border = thin_border
        ws.cell(row=row_idx, column=2).border = thin_border
        ws.cell(row=row_idx, column=3).border = thin_border

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 35

    wb.save(out_path)
    print(f"Acta de Constitución guardada en: {out_path}")

def populate_registro_interesados():
    src_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/MATERIALES/Formato_Registro_de_Interesados.xlsx'
    out_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/Registro_de_Interesados.xlsx'
    
    wb = openpyxl.load_workbook(src_path)
    ws = wb["5 Registro de Interesados"]

    font_content = Font(name='Arial', size=9, bold=False, color='000000')
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    # Lista natural y realista de interesados (sin personas ficticias)
    interesados_data = [
        [
            "Patrocinador (Sponsor)",
            "Dirección de Operaciones y Finanzas",
            "Director de Operaciones",
            "Gerencia General / Finanzas",
            "Que el proyecto se culmine dentro del presupuesto y reduzca los tiempos de cola en tienda.",
            "Alto / Alto",
            "Gestionar atentamente (Reuniones de avance quincenal y reportes de estado).",
            "Sí",
            "-",
            "-"
        ],
        [
            "Gerente del Proyecto (PM)",
            "Max Anderson Benites Corazón",
            "Responsable de Proyecto",
            "Equipo de Proyectos (Estudiante)",
            "Cumplimiento del cronograma, coordinación fluida del equipo y entrega con calidad.",
            "Alto / Alto",
            "Gestionar atentamente (Liderazgo directo y seguimiento continuo a las actividades).",
            "Sí",
            "-",
            "-"
        ],
        [
            "Jefe de Tienda / Operaciones",
            "Administración de Tienda",
            "Gerente de Tienda",
            "Operaciones en Tienda",
            "Que la instalación no perjudique las ventas diarias y que el sistema sea fácil de operar.",
            "Alto / Alto",
            "Gestionar atentamente (Coordinación de horarios de trabajo e información previa).",
            "Sí",
            "-",
            "-"
        ],
        [
            "Área de TI y Sistemas",
            "Jefatura de Soporte y TI",
            "Jefe de TI",
            "Tecnología de la Información",
            "Compatibilidad del software, seguridad en la red y estabilidad en las transacciones.",
            "Alto / Medio",
            "Mantener satisfecho (Validar requerimientos técnicos y arquitectura de red).",
            "Sí",
            "-",
            "-"
        ],
        [
            "Personal de Caja / Supervisores",
            "Equipo de Cajas y Atención",
            "Supervisores y Cajeros",
            "Atención al Cliente",
            "Capacitación clara, facilidad para resolver dudas de clientes y soporte ante fallas.",
            "Bajo / Alto",
            "Mantener informado (Capacitación práctica y entrega de guías de uso).",
            "No",
            "-",
            "-"
        ],
        [
            "Clientes de la Tienda",
            "Público Comprador",
            "Cliente Final",
            "Comunidad de Clientes",
            "Atención rápida, sistema intuitivo y facilidad para pagar sin hacer colas largas.",
            "Bajo / Alto",
            "Monitorear (Observar aceptación en tienda y colocar señalización clara).",
            "No",
            "-",
            "-"
        ],
        [
            "Proveedor de Equipos y Pagos",
            "Empresa Proveedora de Soluciones POS",
            "Ejecutivo de Soporte / Enlace",
            "Proveedor Externo",
            "Entrega oportuna de equipos, especificaciones claras y pagos según contrato.",
            "Medio / Medio",
            "Mantener informado (Coordinar cronograma de entrega y pruebas de enlace).",
            "No",
            "-",
            "-"
        ]
    ]

    start_row = 5
    for i, row_data in enumerate(interesados_data):
        current_row = start_row + i
        ws.row_dimensions[current_row].height = 32
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = val
            cell.font = font_content
            cell.border = thin_border
            if col_idx in [6, 8, 9, 10]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    col_widths = {
        'A': 22,
        'B': 26,
        'C': 24,
        'D': 24,
        'E': 36,
        'F': 18,
        'G': 32,
        'H': 16,
        'I': 8,
        'J': 8
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(out_path)
    print(f"Registro de Interesados guardado en: {out_path}")

def update_zip():
    zip_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/Semana04_Tarea_Acta_Constitucion_y_Registro_Interesados_Max_Benites.zip'
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        zipf.write(
            '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/Acta_de_Constitucion.xlsx',
            arcname='Acta_de_Constitucion.xlsx'
        )
        zipf.write(
            '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/Registro_de_Interesados.xlsx',
            arcname='Registro_de_Interesados.xlsx'
        )
    print(f"ZIP actualizado en: {zip_path}")

if __name__ == '__main__':
    populate_acta_constitucion()
    populate_registro_interesados()
    update_zip()
