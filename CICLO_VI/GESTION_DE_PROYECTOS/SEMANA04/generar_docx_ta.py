import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def build_corrected_document():
    doc = docx.Document()

    # Configurar márgenes estándar APA (2.54 cm / 1 pulgada)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Configurar estilo Normal: Arial 11 pt, interlineado 1.0 (sencillo), espaciado posterior 6 pt
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Arial'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.0
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Propiedades del documento
    core = doc.core_properties
    core.author = 'Grupo de Trabajo'
    core.last_modified_by = 'Alvaro Soto'
    core.title = 'Tarea Académica - Gestión de Proyectos'

    # -------------------------------------------------------------
    # PORTADA OFICIAL UTP
    # -------------------------------------------------------------
    logo_path = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/logo_utp.png'
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(14)
    if os.path.exists(logo_path):
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.8))

    p_fac = doc.add_paragraph()
    p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fac.paragraph_format.space_before = Pt(0)
    p_fac.paragraph_format.space_after = Pt(28)
    r = p_fac.add_run('FACULTAD DE INGENIERÍA')
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True

    p_ta = doc.add_paragraph()
    p_ta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ta.paragraph_format.space_before = Pt(0)
    p_ta.paragraph_format.space_after = Pt(28)
    r = p_ta.add_run('TAREA ACADÉMICA')
    r.font.name = 'Arial'
    r.font.size = Pt(16)
    r.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(22)
    r = p_sub.add_run('Análisis Comparativo de Aplicación Práctica de Metodologías Clásicas (PMI) y Ágiles (Scrum) en Empresas de la Localidad')
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    r.font.italic = True

    p_cur = doc.add_paragraph()
    p_cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cur.paragraph_format.space_before = Pt(0)
    p_cur.paragraph_format.space_after = Pt(2)
    r = p_cur.add_run('CURSO:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    p_curnom = doc.add_paragraph()
    p_curnom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_curnom.paragraph_format.space_before = Pt(0)
    p_curnom.paragraph_format.space_after = Pt(18)
    r = p_curnom.add_run('GESTIÓN DE PROYECTOS (100000I33N) - Sección 13451')
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    r.font.bold = True

    p_doc = doc.add_paragraph()
    p_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_doc.paragraph_format.space_before = Pt(0)
    p_doc.paragraph_format.space_after = Pt(2)
    r = p_doc.add_run('DOCENTE:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    p_docnom = doc.add_paragraph()
    p_docnom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_docnom.paragraph_format.space_before = Pt(0)
    p_docnom.paragraph_format.space_after = Pt(18)
    r = p_docnom.add_run('Nora Noelia Sernaque Zapata')
    r.font.name = 'Arial'
    r.font.size = Pt(11)

    p_int_tit = doc.add_paragraph()
    p_int_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_int_tit.paragraph_format.space_before = Pt(0)
    p_int_tit.paragraph_format.space_after = Pt(6)
    r = p_int_tit.add_run('INTEGRANTES:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    integrantes = [
        'Benites Corazón, Max Anderson – Código: U24217839',
        'Oliva Seminario, Anderson Harold – Código: U21222001',
        'Sandoval Santamaria, Henry Maykol – Código: U20224902',
        'Soto Chavarri, Alvaro Stefano – Código: U23259559'
    ]
    for integ in integrantes:
        p_int = doc.add_paragraph()
        p_int.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_int.paragraph_format.space_before = Pt(0)
        p_int.paragraph_format.space_after = Pt(3)
        r = p_int.add_run(integ)
        r.font.name = 'Arial'
        r.font.size = Pt(11)

    p_pie = doc.add_paragraph()
    p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pie.paragraph_format.space_before = Pt(20)
    p_pie.paragraph_format.space_after = Pt(0)
    r = p_pie.add_run('Lima - Perú\n2026')
    r.font.name = 'Arial'
    r.font.size = Pt(11)

    doc.add_page_break()

    # -------------------------------------------------------------
    # CASO 1: METODOLOGÍAS DE GESTIÓN DE PROYECTOS CLÁSICA (PMI)
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    r = h1.add_run('A. Caso 1: Metodologías de Gestión de Proyectos Clásica')
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True

    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('I. Descripción de la empresa y/o Área afectada por el proyecto a implementar:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    # a) Nombre de la empresa y/o Área
    p = doc.add_paragraph()
    r = p.add_run('a) Nombre de la empresa y/o Área')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'De acuerdo con el registro oficial de la Superintendencia Nacional de Aduanas y de Administración Tributaria (SUNAT, 2026), los datos de identificación corporativa de la entidad en el Perú corresponden a:\n'
        '• Razón Social: N.C.R. DEL PERU S.A.C. (subsidiaria de NCR VOYIX Corporation)\n'
        '• RUC: 20100128137\n'
        '• Domicilio Fiscal: Calle Las Orquídeas Nro. 585, Urb. Jardín (Edificio Fibra), San Isidro, Lima, Perú\n'
        '• Área afectada: Área de Servicios Profesionales e Implementación Técnica (Professional Services & Technical Implementation - Retail Solutions)'
    )

    # b) Descripción de la empresa
    p = doc.add_paragraph()
    r = p.add_run('b) Descripción de la empresa')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Conforme a la documentación corporativa oficial de la compañía (NCR Voyix, 2023), NCR VOYIX Corporation es una empresa global líder en el desarrollo e integración de plataformas de comercio unificado, sistemas de punto de venta (POS) y terminales de autoservicio (Self-Checkout - SCO) para los sectores de retail, restaurantes y banca. Tras la reestructuración estratégica corporativa de NCR Corporation en 2023, NCR VOYIX consolidó de forma exclusiva las soluciones tecnológicas de comercio y canales de autoservicio en tienda.\n\n'
        'En el ámbito local, su filial N.C.R. DEL PERU S.A.C. opera desde su sede en San Isidro (SUNAT, 2026), brindando consultoría técnica, homologación de pasarelas de pago, soporte de ingeniería y despliegue de infraestructura transaccional de misión crítica para las cadenas de hipermercados y supermercados del país.'
    )

    # c) Misión, Visión de la empresa
    p = doc.add_paragraph()
    r = p.add_run('c) Misión, Visión de la empresa')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'De acuerdo con los postulados y lineamientos estratégicos publicados por la compañía (NCR Voyix, 2023):\n'
        '• Misión: Transformar, conectar y optimizar las plataformas de comercio y punto de venta para permitir que los negocios minoristas ofrezcan experiencias de compra fluidas y seguras a los clientes, impulsando la eficiencia operativa mediante soluciones integradas de hardware, software y servicios de despliegue profesional.\n\n'
        '• Visión: Ser el socio estratégico líder a nivel mundial en comercio unificado y autoservicio, redefiniendo las transacciones y la interacción operativa entre las cadenas comerciales y los consumidores finales.'
    )

    # d) Descripción del proyecto
    p = doc.add_paragraph()
    r = p.add_run('d) Descripción del proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '• Nombre del proyecto: Implementación y Rollout Nacional de Terminales de Autoservicio (Self-Checkout - SCO) y Modernización de Plataforma de Punto de Venta (POS) para Cadena de Retail Supermercadista en Lima y Provincias.\n\n'
        '• Objetivo: Descongestionar las líneas de cajas en horarios pico en un 35%, optimizar los costos de atención transaccional y mejorar la experiencia de compra del consumidor mediante estaciones de pago autónomas (NCR Voyix, 2023).\n\n'
        '• Alcance del proyecto:\n'
        '1. Suministro, desaduanaje e importación de 80 terminales de autoservicio NCR SelfServ Checkout (NCR Voyix, 2023).\n'
        '2. Actualización del software base de punto de venta (POS) en 160 estaciones tradicionales en 20 hipermercados a nivel nacional.\n'
        '3. Acondicionamiento de infraestructura en tienda: canalización eléctrica estabilizada, tendido de cableado estructurado categoría 6A certificado y anclaje físico de mobiliario SCO.\n'
        '4. Integración y homologación de periféricos: balanzas de seguridad con detección de peso asistido, escáneres biópticos 2D e integración certificada de pasarelas de pago PinPad (Niubiz/Izipay).\n'
        '5. Despliegue en ventanas nocturnas de corte de tienda (Go-Live windows), capacitación técnica presencial a supervisores y periodo de estabilización post-arranque (Hypercare).'
    )

    # e) Metodología Implementada
    p = doc.add_paragraph()
    r = p.add_run('e) Metodología Implementada')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Se implementó la Metodología Clásica Predictiva basada en los lineamientos del Project Management Institute (PMI, 2017, 2021) y la Guía del PMBOK.\n\n'
        'Justificación: Conforme a la fundamentación del PMI (2017), el enfoque predictivo resulta indispensable cuando el proyecto cuenta con dependencias secuenciales estrictas que conforman una ruta crítica (CPM): no es técnicamente viable instalar el software ni certificar las pasarelas bancarias si previamente no se concluyó la importación, la obra civil, la energía estabilizada y el cableado de red. Asimismo, el proyecto opera bajo contratos con penalidades por indisponibilidad y cronogramas nocturnos cerrados donde la tienda debe abrir a las 8:00 a. m. con operatividad al 100%, lo cual demanda una planificación exhaustiva, control de adquisiciones y planes formales de contingencia (Rollback).'
    )

    # SECCIÓN II
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('II. Aplicación: Identificación de características Generales del modelo PMI:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    # a) Roles y responsabilidades del Project Manager
    p = doc.add_paragraph()
    r = p.add_run('a) Roles y responsabilidades del Project Manager en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Conforme a las áreas de conocimiento y estándares de gestión de la dirección de proyectos (PMI, 2017), el Project Manager asumió la dirección integral del proyecto cumpliendo las siguientes responsabilidades:\n'
        '1. Elaboración de documentos fundacionales: Desarrollo del Acta de Constitución del Proyecto (Project Charter) y la Estructura de Desglose del Trabajo (EDT/WBS), delimitando formalmente el alcance técnico y los paquetes de trabajo (PMI, 2017).\n'
        '2. Planificación y control del cronograma maestro en ruta crítica (MS Project), programando las intervenciones de las 20 tiendas en cuatro fases de despliegue.\n'
        '3. Gestión de la cadena de suministro y adquisiciones, coordinando con plantas internacionales y agencias de aduanas la llegada oportuna del equipamiento a almacenes.\n'
        '4. Coordinación de comunicaciones entre los interesados clave: gerencia de operaciones del cliente, áreas de TI, contratistas de obras civiles y especialistas de campo de NCR.\n'
        '5. Control integrado de cambios mediante solicitudes formales (Change Requests), analizando el impacto presupuestal y temporal de cualquier ajuste de diseño en tienda (PMI, 2017).\n'
        '6. Supervisión de los protocolos de contingencia y planes de reversión (Rollback plan) para asegurar la apertura comercial continua de las tiendas.'
    )

    # b) Habilidades requeridas para un Project Manager
    p = doc.add_paragraph()
    r = p.add_run('b) Habilidades requeridas para un Project Manager en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Conforme al Triángulo de Talentos del PMI (PMI Talent Triangle) definido por el Project Management Institute (PMI, 2021) y estudiado en el marco conceptual del curso, se identificaron tres dimensiones esenciales de competencias:\n\n'
        '1. Formas de trabajar (Gestión Técnica de Proyectos): Dominio de la programación de ruta crítica (CPM), gestión de la EDT, uso de la técnica de Valor Ganado (EVM) para el seguimiento de índices de desempeño de costo (CPI) y cronograma (SPI), y conocimientos técnicos en arquitecturas transaccionales POS, protocolos de comunicación y normas de seguridad PCI-DSS (PMI, 2017, 2021).\n\n'
        '2. Habilidades de poder (Power Skills): Liderazgo para conducir equipos multidisciplinarios (especialistas de software, electricistas y auditores), capacidad de toma de decisiones bajo presión durante ventanas de madrugada y comunicación asertiva para negociar cronogramas con los administradores de tienda sin perturbar las ventas diurnas (PMI, 2021).\n\n'
        '3. Visión empresarial (Business Acumen): Comprensión de los indicadores clave del negocio minorista (costo transaccional por caja, flujo de clientes en horas punta y cumplimiento tributario de comprobantes electrónicos ante SUNAT), alineando la ejecución técnica con las metas estratégicas del cliente (PMI, 2021).'
    )

    # c) Ciclo de Vida del Proyecto
    p = doc.add_paragraph()
    r = p.add_run('c) Ciclo de Vida del Proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'El proyecto siguió un ciclo de vida predictivo estructurado en los cinco grupos de procesos del PMBOK (PMI, 2017):\n\n'
        '1. Inicio: Emisión y aprobación del Acta de Constitución del Proyecto y registro inicial de interesados clave (gerentes de operaciones, líderes de TI, supervisores de caja y proveedores de enlace).\n'
        '2. Planificación: Desarrollo del Plan para la Dirección del Proyecto, descomposición de la EDT, elaboración del cronograma maestro en ruta crítica, diseño de la matriz de riesgos (demoras aduaneras, fallas de red, variaciones de voltaje) y definición de checklists de calidad técnica (PMI, 2017).\n'
        '3. Ejecución: Alistamiento en laboratorio (Staging: clonación de imágenes maestras de sistema operativo y carga de certificados de seguridad), ejecución de obras físicas y eléctricas en tienda, despliegue nocturno por oleadas, pruebas transaccionales de extremo a extremo y capacitación a supervisores de tienda.\n'
        '4. Monitoreo y Control: Medición periódica de avances físicos vs. financieros (EVM), verificación de calidad caja por caja mediante listas de control, gestión de incidencias en etapa de estabilización (Hypercare de 7 días) y tramitación formal de solicitudes de cambio (PMI, 2017).\n'
        '5. Cierre: Suscripción de actas de conformidad técnica (Sign-off) con cada gerente de tienda, transferencia de la solución a la mesa de ayuda y soporte 24/7 de NCR VOYIX, cierre de contratos de servicios y documentación de lecciones aprendidas (PMI, 2017).'
    )

    # d) Beneficios del uso de la metodología
    p = doc.add_paragraph()
    r = p.add_run('d) Beneficios del uso de la metodología en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'De acuerdo con los principios de desempeño predictivo documentados por el PMI (2017, 2021), los beneficios obtenidos fueron:\n'
        '• Previsibilidad en plazos y presupuesto: Al establecer una línea base detallada, se controlaron estrictamente los costos de adquisición de hardware internacional y horas de trabajo de campo.\n'
        '• Control riguroso de dependencias físicas y logísticas: Sincronizó con exactitud la llegada de equipos importados, obras civiles y conectividad antes del ingreso de los técnicos a tienda.\n'
        '• Continuidad operativa asegurada: La planificación nocturna y los planes de Rollback garantizaron que la apertura de tiendas y las ventas al público no sufrieran interrupciones.\n'
        '• Estandarización y respaldo auditable: La documentación formal del PMBOK permitió una trazabilidad completa del proyecto ante auditorías internas y entidades financieras.'
    )

    # e) Desventajas del uso de la metodología
    p = doc.add_paragraph()
    r = p.add_run('e) Desventajas del uso de la metodología en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Conforme a las limitaciones inherentes a los modelos predictivos señaladas por el PMI (2017, 2021), se presentaron las siguientes desventajas:\n'
        '• Baja adaptabilidad ante cambios de último momento: Modificaciones no programadas en la distribución física de pasillos requirieron trámites formales de cambio que dilataron los tiempos de respuesta.\n'
        '• Exigente carga administrativa: La elaboración y firma continua de actas, minutas y reportes demandó tiempo de gestión significativo para el equipo de proyectos.\n'
        '• Percepción tardía del valor: A diferencia de los modelos iterativos, el cliente solo experimenta el beneficio directo de la solución cuando la tienda completa íntegramente todo el ciclo de instalación física y certificación nocturna.'
    )

    # SECCIÓN III: CONCLUSIONES CASO CLÁSICO
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('III. Conclusiones.')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '1. La metodología clásica PMI/PMBOK resulta el enfoque idóneo para proyectos de implementación de infraestructura física transaccional y rollout masivo en retail, dado que la presencia de dependencias secuenciales inflexibles (aduanas, adecuación eléctrica y de red) requiere un control riguroso de la ruta crítica que las metodologías adaptativas no están diseñadas para gestionar en entornos de tienda activa.\n\n'
        '2. El liderazgo del Project Manager bajo el estándar PMI trasciende la supervisión administrativa al integrar armónicamente las habilidades técnicas de control (EVM, EDT, CPM) con habilidades interpersonales y visión empresarial, permitiendo articular proveedores logísticos, contratistas y gerencias de tienda para garantizar el Go-Live nocturno sin riesgo de interrupción de ventas comerciales.'
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # CASO 2: METODOLOGÍAS DE GESTIÓN DE PROYECTOS ÁGILES (SCRUM)
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    r = h1.add_run('B. Caso 2: Metodologías de Gestión de Proyectos Ágiles')
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True

    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('IV. Descripción de la empresa y/o Área afectada por el proyecto a implementar:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    # f) Nombre de la empresa y/o Área
    p = doc.add_paragraph()
    r = p.add_run('f) Nombre de la empresa y/o Área')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'De acuerdo con los registros oficiales de la Superintendencia Nacional de Aduanas y de Administración Tributaria (SUNAT, 2026) e información corporativa de la entidad, los datos de identificación corresponden a:\n'
        '• Razón Social: BANCO DE CREDITO DEL PERU (subsidiaria de Credicorp Ltd.)\n'
        '• RUC: 20100047218\n'
        '• Domicilio Fiscal: Calle Centenario Nro. 156, Urb. Las Laderas de Melgarejo, La Molina, Lima, Perú\n'
        '• Producto: Yape (Billetera Digital)\n'
        '• Área afectada: Centro de Innovación (CEX) y Tribu de Préstamos Digitales de Yape'
    )

    # g) Descripción de la empresa
    p = doc.add_paragraph()
    r = p.add_run('g) Descripción de la empresa')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'El Banco de Crédito del Perú (BCP) es la institución financiera más grande y antigua del país. Dentro de su proceso de transformación digital, el BCP desarrolló Yape, una billetera electrónica que nació como una iniciativa intraemprendedora. Hoy en día, Yape es el principal aplicativo de pagos móviles en el Perú, con millones de usuarios activos, que permite realizar transferencias con el número de celular o código QR, pago de servicios, recargas y acceso a productos financieros rápidos (BCP, 2023).'
    )

    # h) Misión, Visión de la empresa
    p = doc.add_paragraph()
    r = p.add_run('h) Misión, Visión de la empresa')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Según los lineamientos de responsabilidad social e inclusión financiera del BCP y Yape (BCP, 2023):\n'
        '• Misión: Impulsar la inclusión financiera en el Perú democratizando el acceso a los medios de pago y servicios financieros digitales, ofreciendo una experiencia de usuario rápida, segura y gratuita.\n\n'
        '• Visión: Ser el ecosistema digital transaccional y financiero preferido por todos los peruanos, integrando soluciones de pago, crédito y comercio en una sola plataforma de uso diario.'
    )

    # i) Descripción del proyecto
    p = doc.add_paragraph()
    r = p.add_run('i) Descripción del proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '• Nombre del proyecto: Desarrollo y Despliegue de la Funcionalidad "Microcréditos Yape" (Préstamos al instante).\n\n'
        '• Objetivo: Diseñar, desarrollar y lanzar un módulo dentro del aplicativo Yape que permita a los usuarios precalificados solicitar y recibir microcréditos entre S/ 50 y S/ 500 con desembolso automático en menos de 5 segundos.\n\n'
        '• Alcance del proyecto:\n'
        '1. Desarrollo del front-end UI/UX del flujo de solicitud de préstamo en la app móvil.\n'
        '2. Integración back-end con el motor de riesgo crediticio del BCP para validación en tiempo real.\n'
        '3. Pruebas A/B con un grupo cerrado de usuarios beta testers para medir la usabilidad.\n'
        '4. Salida a producción escalonada (Rollout) al mercado nacional.'
    )

    # j) Metodología Implementada
    p = doc.add_paragraph()
    r = p.add_run('j) Metodología Implementada')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Se implementó la Metodología Ágil basada en el marco de trabajo Scrum, fundamentada en los principios del Manifiesto Ágil y la Guía de Scrum (Schwaber & Sutherland, 2020).\n\n'
        'Justificación: A diferencia de la instalación de hardware en tiendas retail, el desarrollo de un producto de software financiero enfrenta alta incertidumbre, requerimientos cambiantes y la necesidad de probar la respuesta del mercado rápidamente (Time-to-Market). Scrum permite trabajar mediante iteraciones cortas, obteniendo feedback real de los usuarios para mejorar la funcionalidad del crédito antes de liberarla a millones de personas, minimizando así los riesgos de adopción y fallos masivos en producción.'
    )

    # SECCIÓN V
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('V. Aplicación: Identificación de características Generales del modelo Ágil:')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    # f) Roles y responsabilidades del Scrum Master
    p = doc.add_paragraph()
    r = p.add_run('f) Roles y responsabilidades del Scrum Master en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'Siguiendo los lineamientos de la Guía de Scrum (Schwaber & Sutherland, 2020), el Scrum Master del proyecto de "Microcréditos Yape" desempeñó las siguientes responsabilidades:\n'
        '1. Facilitador de ceremonias: Garantizar que las reuniones de Sprint Planning, Daily Scrum, Sprint Review y Sprint Retrospective se realicen de manera productiva y dentro del tiempo establecido (timebox).\n'
        '2. Removedor de impedimentos: Solucionar bloqueos operativos, como demoras en los accesos a las APIs del motor de riesgo del BCP, para que los desarrolladores no detengan su trabajo.\n'
        '3. Agente de cambio y mentor: Guiar al Product Owner en técnicas efectivas para gestionar y priorizar el Product Backlog, asegurando que los requerimientos de cumplimiento normativo de la SBS estén claros.\n'
        '4. Protector del equipo: Aislar al equipo de desarrollo de interrupciones externas (como solicitudes urgentes de otras gerencias durante el desarrollo del Sprint).'
    )

    # g) Habilidades requeridas para un Scrum Master
    p = doc.add_paragraph()
    r = p.add_run('g) Habilidades requeridas para un Scrum Master en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '1. Liderazgo Servicial: Liderar desde el apoyo y la facilitación, empoderando al equipo de desarrolladores en lugar de ejercer un mando y control jerárquico tradicional.\n'
        '2. Resolución de Conflictos: Habilidad para mediar en desacuerdos técnicos entre los arquitectos de software y el área de ciberseguridad financiera.\n'
        '3. Comunicación Ágil: Capacidad para transmitir la mentalidad ágil (Agile Mindset) y asegurar que el equipo comprenda el valor del producto mínimo viable (MVP).\n'
        '4. Adaptabilidad y Flexibilidad: Capacidad para ajustar las métricas y tableros (Jira, Trello) para reflejar transparentemente el progreso en un entorno de trabajo altamente dinámico.'
    )

    # h) Ciclo de Vida del Proyecto
    p = doc.add_paragraph()
    r = p.add_run('h) Ciclo de Vida del Proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        'El ciclo de vida ágil se ejecutó de forma iterativa e incremental en iteraciones (Sprints) de 2 semanas:\n'
        '1. Visión y Backlog: El Product Owner definió la visión del módulo de préstamos y creó un Product Backlog priorizado con Historias de Usuario (por ejemplo: "Como usuario precalificado, quiero ver el monto máximo que me pueden prestar para tomar una decisión informada").\n'
        '2. Sprint Planning: El equipo de desarrollo seleccionó las Historias de Usuario con mayor prioridad y elaboró el Sprint Backlog para las siguientes dos semanas.\n'
        '3. Ejecución y Daily Scrum: El equipo programó y probó el código diariamente. Cada mañana realizaron una reunión de 15 minutos para sincronizar actividades y reportar bloqueos (Daily Stand-up).\n'
        '4. Sprint Review: Al final del Sprint, el equipo presentó un "Incremento" de software funcional (la pantalla de solicitud simulada y validación de APIs) a los stakeholders (Gerencia del BCP y área de Riesgos) para recibir retroalimentación.\n'
        '5. Sprint Retrospective: Tras la revisión, el equipo analizó sus procesos internos para identificar oportunidades de mejora técnica y humana para el siguiente Sprint.'
    )

    # i) Beneficios del uso de la metodología
    p = doc.add_paragraph()
    r = p.add_run('i) Beneficios del uso de la metodología en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '• Rápida entrega de valor: Permitió lanzar una versión básica (MVP) de los préstamos en pocos meses, generando ingresos y validando el interés del usuario rápidamente.\n'
        '• Adaptabilidad al cambio: Fue posible ajustar la tasa de interés y los límites de monto prestado en función del feedback de los primeros usuarios sin desestabilizar todo el proyecto.\n'
        '• Mitigación de riesgos financieros: Al liberar el producto en ciclos pequeños, se detectaron y corrigieron tempranamente vulnerabilidades de seguridad antes de una exposición nacional.\n'
        '• Alta colaboración: Fomentó una sinergia constante entre el equipo de tecnología y las áreas de negocio y cumplimiento normativo del banco.'
    )

    # j) Desventajas del uso de la metodología
    p = doc.add_paragraph()
    r = p.add_run('j) Desventajas del uso de la metodología en el proyecto')
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '• Incertidumbre en el costo y alcance final: Al no tener una línea base estricta desde el inicio, resulta difícil predecir para la gerencia financiera el costo total exacto del desarrollo de la funcionalidad.\n'
        '• Alta dependencia del compromiso del equipo: Scrum requiere equipos autoorganizados y altamente maduros; si un desarrollador clave falla o rota, el impacto en la velocidad del Sprint es severo.\n'
        '• Desgaste por intensidad: El ritmo constante de entregas cada dos semanas puede generar agotamiento en el equipo si no se gestiona bien la capacidad y la deuda técnica.'
    )

    # SECCIÓN VI: CONCLUSIONES
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('VI. Conclusiones.')
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        '1. Idoneidad del modelo predictivo en despliegues físicos (Caso NCR VOYIX): La metodología clásica PMI/PMBOK resulta el enfoque idóneo para proyectos de implementación de infraestructura física transaccional y rollout masivo en retail, dado que la presencia de dependencias secuenciales inflexibles (aduanas, adecuación eléctrica y cableado de red) requiere un control riguroso de la ruta crítica que las metodologías adaptativas no están diseñadas para gestionar en entornos de tienda activa.\n\n'
        '2. Agilidad y reducción del time-to-market en productos fintech (Caso BCP - Yape): La metodología ágil Scrum demostró ser altamente efectiva para el desarrollo de soluciones de software financiero como la funcionalidad de "Microcréditos Yape", debido a que las iteraciones cortas de dos semanas y la validación continua con usuarios reales permitieron calibrar las reglas de riesgo y mejorar la experiencia de usuario antes del lanzamiento masivo nacional, reduciendo sustancialmente los riesgos de fallas en producción.\n\n'
        '3. Liderazgo situacional articulador y facilitador: En ambos modelos el factor humano y el liderazgo resultan determinantes para el éxito del proyecto: mientras que bajo el estándar PMI el Project Manager articula contratistas de obra, proveedores logísticos y gerencias mediante herramientas de control formal (EVM, EDT, CPM) asegurando ventanas nocturnas de corte sin paradas operativas, en el marco Scrum el Scrum Master actúa como líder servicial que remueve bloqueos técnicos en APIs corporativas y protege al squad de presiones externas para entregar incrementos de software estables.\n\n'
        '4. Complementariedad metodológica en la gestión corporativa: Al contrastar los proyectos estudiados, se determina que no existe una metodología universalmente superior. El enfoque clásico PMI es indispensable para proyectos con restricciones físicas, logísticas y alta necesidad de predictibilidad de costos, mientras que los modelos ágiles destacan en entornos de alta incertidumbre y desarrollo de productos digitales donde la capacidad de adaptación rápida a la retroalimentación del mercado constituye el principal factor de éxito.'
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # REFERENCIAS BIBLIOGRÁFICAS (NORMAS APA 7.ª EDICIÓN)
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(10)
    r = h1.add_run('Referencias Bibliográficas')
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True

    referencias = [
        'Banco de Crédito del Perú [BCP]. (2023). Reporte de Sostenibilidad e Inclusión Financiera BCP. Credicorp Ltd. https://www.viabcp.com',
        'NCR Voyix. (2023). Unified commerce and self-checkout solutions for retailers: Corporate overview. NCR Voyix Corporation. https://www.ncrvoyix.com',
        'Project Management Institute. (2017). Guía de los fundamentos para la dirección de proyectos (Guía del PMBOK) (6.ª ed.). Project Management Institute.',
        'Project Management Institute. (2021). Guía de los fundamentos para la dirección de proyectos (Guía del PMBOK) y el estándar para la dirección de proyectos (7.ª ed.). Project Management Institute.',
        'Schwaber, K., & Sutherland, J. (2020). La Guía de Scrum: Las reglas del juego. Scrum.org. https://scrumguides.org',
        'Superintendencia Nacional de Aduanas y de Administración Tributaria. (2026). Consulta RUC: BANCO DE CREDITO DEL PERU (RUC 20100047218). Plataforma del Estado Peruano. https://e-consultaruc.sunat.gob.pe',
        'Superintendencia Nacional de Aduanas y de Administración Tributaria. (2026). Consulta RUC: N.C.R. DEL PERU S.A.C. (RUC 20100128137). Plataforma del Estado Peruano. https://e-consultaruc.sunat.gob.pe'
    ]

    for ref in referencias:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.0
        r = p_ref.add_run(ref)
        r.font.name = 'Arial'
        r.font.size = Pt(11)

    # Guardar en la carpeta de avance del compañero y en la carpeta principal
    output_path1 = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/TA_AVANCE_COMPAÑEROS/TA_Gestion_de_Proyectos.docx'
    output_path2 = '/home/ilkay/Documentos/UTP/CICLO_VI/GESTION_DE_PROYECTOS/SEMANA04/TA_Gestion_de_Proyectos.docx'
    
    doc.save(output_path1)
    doc.save(output_path2)
    print(f'Documento guardado exitosamente en:\n  - {output_path1}\n  - {output_path2}')

if __name__ == '__main__':
    build_corrected_document()
